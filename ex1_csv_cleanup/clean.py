import pandas as pd, re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
df=pd.read_csv("messy_sales.csv",dtype=str)
df.columns=["invoice","customer","region","date","amount","qty"]
before=len(df)
df=df.drop_duplicates()
df["customer"]=df.customer.str.strip().str.title().str.replace("Llc","LLC")
df["region"]=df.region.str.strip().str.title()
df["date"]=pd.to_datetime(df.date.str.strip(),format="mixed").dt.date
df["amount"]=df.amount.str.replace(r"[^\d.]","",regex=True).astype(float)
df["qty"]=pd.to_numeric(df.qty.str.strip(),errors="coerce")
df["qty_missing"]=df.qty.isna()
df["qty"]=df.qty.fillna(1).astype(int)
df=df.sort_values(["date","invoice"])
with pd.ExcelWriter("clean_sales.xlsx",engine="openpyxl") as xw:
    df.to_excel(xw,sheet_name="Data",index=False)
    pv=df.pivot_table(index="customer",columns="region",values="amount",aggfunc="sum",fill_value=0).round(2)
    pv.to_excel(xw,sheet_name="Pivot_Customer_Region")
    pd.DataFrame({"metric":["rows in","duplicates removed","rows out","qty missing (set to 1)"],
                  "value":[before,before-len(df),len(df),int(df.qty_missing.sum())]}).to_excel(xw,sheet_name="Change_Log",index=False)
wb=load_workbook("clean_sales.xlsx"); ws=wb["Data"]; n=ws.max_row
ws["H1"]="line_total"
for r in range(2,n+1): ws[f"H{r}"]=f"=E{r}*F{r}"
ws["J1"]="Total revenue"; ws["K1"]=f"=SUM(H2:H{n})"
ws["J2"]="Avg invoice"; ws["K2"]=f"=AVERAGE(E2:E{n})"
ws["J3"]="Acme revenue"; ws["K3"]=f'=SUMIF(B2:B{n},"Acme Corp",H2:H{n})'
for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="DDEBF7")
ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:H{n}"
for col,w in zip("ABCDEFGH",[12,16,10,12,10,6,12,12]): ws.column_dimensions[col].width=w
wb.save("clean_sales.xlsx"); print("ok",before,len(df))
